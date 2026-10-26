#!/usr/bin/env python3
"""Convert docs/PROJECT_REPORT.md into docs/Project_Report.docx.

    pip install python-docx
    python scripts/build_report_docx.py

Supports the Markdown used in the report: headings, paragraphs, bullet and
numbered lists, tables, fenced code blocks, images and inline **bold**,
*italic*, `code` and [links](url).
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parent.parent / "docs"
SRC = DOCS / "PROJECT_REPORT.md"
OUT = DOCS / "Project_Report.docx"

ACCENT = RGBColor(0x43, 0x38, 0xCA)
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|_[^_\s][^_]*_|``[^`]+``|`[^`]+`|\[[^\]]+\]\([^)]+\)|<https?://[^>]+>)")


def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_inline(paragraph, text: str, bold: bool = False, size: float | None = None) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        run_bold, italic, mono = bold, False, False
        if part.startswith("**") and part.endswith("**"):
            part, run_bold = part[2:-2], True
            if part.startswith("`") and part.endswith("`"):
                part, mono = part.strip("`"), True
        elif part.startswith("``") and part.endswith("``"):
            part, mono = part[2:-2].strip(), True
        elif part.startswith("`") and part.endswith("`"):
            part, mono = part[1:-1], True
        elif (part.startswith("*") and part.endswith("*")) or (part.startswith("_") and part.endswith("_")):
            part, italic = part[1:-1], True
        elif part.startswith("[") and "](" in part:
            label, url = part[1:-1].split("](", 1)
            part = label if url.startswith("http") is False else f"{label} ({url})"
        elif part.startswith("<http"):
            part = part[1:-1]
        run = paragraph.add_run(part)
        run.bold = run_bold
        run.italic = italic
        if mono:
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x37, 0x30, 0xA3)
        elif size:
            run.font.size = Pt(size)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    header, body = rows[0], rows[1:]
    # image grids (rows of ![...](...)) become borderless picture tables
    if all(not h.strip() for h in header) and body and all("![" in c for r in body for c in r if c.strip()):
        table = doc.add_table(rows=len(body), cols=len(header))
        for i, r in enumerate(body):
            for j, c in enumerate(r):
                m = re.search(r"!\[([^\]]*)\]\(([^)]+)\)", c)
                if m and (DOCS / m.group(2)).exists():
                    p = table.cell(i, j).paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(DOCS / m.group(2)), width=Cm(7.8))
                    cap = table.cell(i, j).add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r_ = cap.add_run(m.group(1))
                    r_.italic = True
                    r_.font.size = Pt(8.5)
        return
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        cell = table.rows[0].cells[j]
        shade(cell, "E0E7FF")
        add_inline(cell.paragraphs[0], h.strip(), bold=True, size=9)
    for r in body:
        cells = table.add_row().cells
        for j in range(len(header)):
            add_inline(cells[j].paragraphs[0], (r[j] if j < len(r) else "").strip(), size=9)
    doc.add_paragraph()


def split_row(line: str) -> list[str]:
    line = line.strip().strip("|")
    return [c.replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]


def build() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for lvl, size in ((1, 18), (2, 15), (3, 12.5)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.color.rgb = ACCENT
        st.font.size = Pt(size)

    lines = SRC.read_text(encoding="utf-8").splitlines()

    # ---- cover page from the first two headings
    title = lines[0].lstrip("# ").strip()
    subtitle = lines[1].lstrip("# ").strip()
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(30)
    r.font.color.rgb = ACCENT
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.size = Pt(15)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("React · FastAPI · MySQL 8")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    i = 2
    para: list[str] = []
    in_list = False

    def flush() -> None:
        if para:
            add_inline(doc.add_paragraph(), " ".join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            table = doc.add_table(rows=1, cols=1)
            cell = table.rows[0].cells[0]
            shade(cell, "F1F5F9")
            p = cell.paragraphs[0]
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            doc.add_paragraph()
        elif s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            text = s.lstrip("#").strip()
            if level == 2 and text[:1].isdigit() is False and text not in ("Abstract", "References"):
                level = 2
            doc.add_heading(text, level=min(level - 1, 3) if level > 1 else 1)
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\|(\s*:?-+:?\s*\|)+$", lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        elif re.match(r"^[*-] ", s):
            flush()
            add_inline(doc.add_paragraph(style="List Bullet"), s[2:])
            in_list = True
            i += 1
            continue
        elif re.match(r"^\d+\. ", s):
            flush()
            add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", s))
            in_list = True
            i += 1
            continue
        elif s == "---":
            flush()
        elif not s:
            flush()
        elif in_list and line.startswith("  "):
            # continuation of a list item
            add_inline(doc.paragraphs[-1], " " + s)
            i += 1
            continue
        else:
            para.append(s)
        in_list = False
        i += 1
    flush()
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
